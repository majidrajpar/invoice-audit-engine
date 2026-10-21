"""
Test Suite for Invoice Audit Engine.
"""

import os
from invoice_audit_engine.pipeline import process_document
from invoice_audit_engine.checkpoint_db import CheckpointLedger
from invoice_audit_engine.excel_exporter import export_audit_to_excel

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
SAMPLE_INVOICE = os.path.join(FIXTURES_DIR, "sample_invoice.png")
SAMPLE_LQ = os.path.join(FIXTURES_DIR, "sample_low_quality.jpg")

def test_clean_invoice_pipeline():
    assert os.path.exists(SAMPLE_INVOICE), "Fixture sample_invoice.png missing"
    res = process_document(SAMPLE_INVOICE)
    assert res.quality_passed is True
    assert res.header is not None
    assert res.header.invoice_number == "US-001"
    assert res.totals is not None
    assert res.totals.subtotal == 145.0
    assert res.totals.total_amount == 154.06
    assert res.audit_flags.is_mathematically_sound is True
    assert res.document_type == "INVOICE"

def test_quality_gate_rejection():
    assert os.path.exists(SAMPLE_LQ), "Fixture sample_low_quality.jpg missing"
    res = process_document(SAMPLE_LQ)
    assert res.quality_passed is False
    assert any("REJECTED" in n for n in res.audit_flags.notes)

def test_checkpoint_ledger_persistence(tmp_path):
    db_file = str(tmp_path / "test_ledger.db")
    ledger = CheckpointLedger(db_file)
    assert len(ledger.get_processed_hashes()) == 0

    res = process_document(SAMPLE_INVOICE)
    ledger.save_result(res)
    hashes = ledger.get_processed_hashes()
    assert res.file_hash_sha256 in hashes

    results = ledger.load_all_results()
    assert len(results) == 1
    assert results[0].header.invoice_number == "US-001"

def test_excel_export_multi_tab(tmp_path):
    res = process_document(SAMPLE_INVOICE)
    out_file = str(tmp_path / "test_report.xlsx")
    saved = export_audit_to_excel([res], out_file)
    assert os.path.exists(saved)
    assert os.path.getsize(saved) > 5000

def test_amount_cleaner_heuristics():
    from invoice_audit_engine.spatial_parser import clean_amount_str
    # Standard amounts
    assert clean_amount_str("100.00") == 100.0
    assert clean_amount_str("1,250.50") == 1250.5
    assert clean_amount_str("AED 500.00") == 500.0
    # Accounting negatives
    assert clean_amount_str("(25.00)") == -25.0
    assert clean_amount_str("-15.50") == -15.5
    assert clean_amount_str("30.00 CR") == -30.0
    # Non-monetary patterns rejected
    assert clean_amount_str("65W") is None
    assert clean_amount_str("CAT-6") is None
    assert clean_amount_str("M8x20") is None
    assert clean_amount_str("15%") is None
    assert clean_amount_str("2026-10-12") is None
    assert clean_amount_str("12/10/2026") is None

def test_line_item_extraction_details():
    res = process_document(SAMPLE_INVOICE)
    assert len(res.line_items) == 3
    
    # Item 1
    assert "Front and rearbrake" in res.line_items[0].description
    assert res.line_items[0].quantity == 1.0
    assert res.line_items[0].unit_price == 100.0
    assert res.line_items[0].line_total == 100.0
    assert res.line_items[0].is_math_sound is True

    # Item 2
    assert "pedalarms" in res.line_items[1].description
    assert res.line_items[1].quantity == 2.0
    assert res.line_items[1].unit_price == 15.0
    assert res.line_items[1].line_total == 30.0
    assert res.line_items[1].is_math_sound is True

    # Item 3
    assert "Labor" in res.line_items[2].description
    assert res.line_items[2].quantity == 3.0
    assert res.line_items[2].unit_price == 5.0
    assert res.line_items[2].line_total == 15.0
    assert res.line_items[2].is_math_sound is True

def test_deskew_angle_accuracy():
    import numpy as np
    import cv2
    from invoice_audit_engine.preprocessor import deskew_image

    h, w = 600, 800
    img = np.full((h, w, 3), 255, dtype=np.uint8)
    for y in range(150, 450, 40):
        cv2.putText(img, "TAX INVOICE AUDIT TEST LINE ITEM 100.00", (100, y), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 0), 2)

    # Rotate by -4.0 degrees
    M = cv2.getRotationMatrix2D((w // 2, h // 2), -4.0, 1.0)
    skewed = cv2.warpAffine(img, M, (w, h), borderValue=(255, 255, 255))

    _, angle = deskew_image(skewed)
    assert abs(angle - 4.0) < 0.5

def test_sqlite_wal_mode_and_concurrency(tmp_path):
    db_file = str(tmp_path / "wal_test.db")
    ledger = CheckpointLedger(db_file)
    with ledger._get_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("PRAGMA journal_mode;")
        mode = cursor.fetchone()[0]
        assert mode.upper() == "WAL"

def test_narrow_thermal_receipt_quality_gate():
    import numpy as np
    from invoice_audit_engine.quality_gate import evaluate_image_quality

    # Create synthetic narrow thermal receipt: width=350, height=900 (aspect ratio 2.57)
    img = np.full((900, 350, 3), 255, dtype=np.uint8)
    # Add text-like high frequency variation so sharpness/contrast pass
    img[100:800:30, 50:300] = 0

    report = evaluate_image_quality(img)
    # Ensure it did NOT fail resolution
    assert not any("Resolution too low" in issue for issue in report.issues)

def test_trailing_minus_and_european_dates():
    from invoice_audit_engine.spatial_parser import clean_amount_str
    import re

    assert clean_amount_str("50.00-") == -50.0
    assert clean_amount_str("1,200.50-") == -1200.5

    date_pattern = re.compile(r"\b(\d{1,4}[/.-]\d{1,2}[/.-]\d{2,4}|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})\b")
    m = date_pattern.search("Invoice Date: 15.01.2024")
    assert m is not None
    assert m.group(1) == "15.01.2024"

def test_document_archetype_classification_and_exclusion(tmp_path):
    from invoice_audit_engine.spatial_parser import SpatialInvoiceParser, OCRToken
    from invoice_audit_engine.schema import InvoiceExtractionResult, AuditValidationFlags, InvoiceHeader, InvoiceTotals

    # 1. Statement of Account
    soa_tokens = [
        OCRToken(text="STATEMENT OF ACCOUNT", xmin=100, ymin=50, xmax=400, ymax=80, confidence=0.99),
        OCRToken(text="Account: ACME Corp", xmin=100, ymin=100, xmax=300, ymax=120, confidence=0.95),
        OCRToken(text="OPENING BALANCE", xmin=100, ymin=200, xmax=250, ymax=220, confidence=0.95),
        OCRToken(text="15,000.00", xmin=400, ymin=200, xmax=500, ymax=220, confidence=0.95),
        OCRToken(text="CLOSING BALANCE", xmin=100, ymin=300, xmax=250, ymax=320, confidence=0.95),
        OCRToken(text="12,500.00", xmin=400, ymin=300, xmax=500, ymax=320, confidence=0.95),
    ]
    parser_soa = SpatialInvoiceParser(soa_tokens)
    assert parser_soa.document_type == "STATEMENT_OF_ACCOUNT"
    assert parser_soa.is_invoice_archetype is False
    assert parser_soa.extract_totals().total_amount is None
    assert len(parser_soa.extract_line_items()) == 0
    flags_soa = parser_soa.reconcile_and_audit(
        parser_soa.extract_header(),
        parser_soa.extract_totals(),
        parser_soa.extract_line_items()
    )
    assert flags_soa.is_mathematically_sound is True
    assert any("EXCLUDED_FROM_AUDIT" in note for note in flags_soa.notes)

    # 2. Delivery Note
    dn_tokens = [
        OCRToken(text="DELIVERY NOTE", xmin=100, ymin=50, xmax=350, ymax=80, confidence=0.98),
        OCRToken(text="Delivered To: Warehouse 4", xmin=100, ymin=100, xmax=350, ymax=120, confidence=0.95),
        OCRToken(text="Item: Steel Pipes Qty: 50", xmin=100, ymin=150, xmax=350, ymax=170, confidence=0.95),
    ]
    parser_dn = SpatialInvoiceParser(dn_tokens)
    assert parser_dn.document_type == "DELIVERY_NOTE"
    assert parser_dn.is_invoice_archetype is False

    # 3. Excel Export handling of excluded documents
    res_inv = process_document(SAMPLE_INVOICE)
    res_soa = InvoiceExtractionResult(
        file_path="C:/invoices/SOA_June2024.pdf",
        file_hash_sha256="fakehashsoa123456789",
        document_type=parser_soa.document_type,
        is_invoice_archetype=parser_soa.is_invoice_archetype,
        exclusion_reason=parser_soa.exclusion_reason,
        quality_passed=True,
        quality_metrics={"issues": []},
        header=parser_soa.extract_header(),
        line_items=[],
        totals=parser_soa.extract_totals(),
        audit_flags=flags_soa,
        raw_token_count=len(soa_tokens),
        average_ocr_confidence=0.95
    )

    out_file = str(tmp_path / "test_exclusion_report.xlsx")
    export_audit_to_excel([res_inv, res_soa], out_file)
    assert os.path.exists(out_file)

    import openpyxl
    wb = openpyxl.load_workbook(out_file)
    ws_sum = wb["Executive_Summary"]
    
    # Check KPI metrics in Tab 1
    kpis = {ws_sum.cell(row=r, column=1).value: ws_sum.cell(row=r, column=2).value for r in range(5, 15)}
    assert kpis["Total Documents Examined"] == 2
    assert kpis["Clean Invoices Audited"] == 1
    assert kpis["Non-Invoice Documents Filtered & Excluded"] == 1
    # Spend should only be the invoice total (154.06), not inflated by SOA
    assert round(kpis["Total Billed Spend Audited"], 2) == 154.06

    # Tab 2: Invoice Register status checks
    ws_reg = wb["Invoice_Register"]
    assert ws_reg.cell(row=2, column=2).value == "PASSED"
    assert ws_reg.cell(row=3, column=2).value == "EXCLUDED (STATEMENT_OF_ACCOUNT)"
    assert ws_reg.cell(row=3, column=16).value == "N/A (Excluded)"

def test_shadow_and_illumination_normalization():
    import numpy as np
    import cv2
    from invoice_audit_engine.preprocessor import normalize_illumination_and_shadows

    # Create synthetic wrinkled paper with a heavy diagonal shadow gradient
    h, w = 300, 400
    base = np.ones((h, w), dtype=np.float32) * 230.0
    # Add dark shadow vignette across bottom-right
    for y in range(h):
        for x in range(w):
            base[y, x] -= ((y / h) * 100.0 + (x / w) * 60.0)

    base_u8 = np.clip(base, 0, 255).astype(np.uint8)
    cv2.putText(base_u8, "TOTAL 125.00", (50, 150), cv2.FONT_HERSHEY_SIMPLEX, 0.8, 20, 2)
    bgr = cv2.cvtColor(base_u8, cv2.COLOR_GRAY2BGR)

    normalized = normalize_illumination_and_shadows(bgr, force=True)
    assert normalized.shape == bgr.shape
    # Check that corner gray variance was substantially flattened towards bright background
    gray_norm = cv2.cvtColor(normalized, cv2.COLOR_BGR2GRAY)
    bg_sample_top_left = gray_norm[10, 10]
    bg_sample_bottom_right = gray_norm[280, 380]
    assert abs(int(bg_sample_top_left) - int(bg_sample_bottom_right)) < 35

def test_multipage_pdf_ingestion(tmp_path):
    import pymupdf
    from invoice_audit_engine.pipeline import process_document

    pdf_path = str(tmp_path / "multipage_invoice.pdf")
    doc = pymupdf.open()
    # Page 1: Header + First Item
    page1 = doc.new_page(width=595, height=842)
    page1.insert_text((50, 50), "Veritux Consulting Services LLC", fontsize=16)
    page1.insert_text((50, 80), "TAX INVOICE # MP-1001", fontsize=14)
    page1.insert_text((50, 110), "Date: 15/01/2026  PO: PO-9988", fontsize=12)
    page1.insert_text((50, 150), "ITEM DESCRIPTION  QTY  RATE  TOTAL", fontsize=12)
    page1.insert_text((50, 190), "Consulting Services  10  50.00  500.00", fontsize=12)
    # Page 2: Second Item + Summary Totals
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((50, 50), "Veritux Consulting Services LLC - Billing Schedule Continued", fontsize=12)
    page2.insert_text((50, 100), "ITEM DESCRIPTION  QTY  RATE  TOTAL", fontsize=12)
    page2.insert_text((50, 140), "Audit Implementation  5  100.00  500.00", fontsize=12)
    page2.insert_text((50, 240), "SUBTOTAL: 1000.00", fontsize=14)
    page2.insert_text((50, 280), "TOTAL AMOUNT: 1000.00", fontsize=16)
    doc.save(pdf_path)
    doc.close()

    res = process_document(pdf_path)
    assert res.quality_passed is True
    assert res.header.invoice_number == "MP-1001"
    assert res.totals.total_amount == 1000.00
    assert len(res.line_items) == 2
    assert res.audit_flags.is_mathematically_sound is True

def test_column_gutter_isolation():
    from invoice_audit_engine.spatial_parser import SpatialInvoiceParser, OCRToken

    # Header line with explicit columns
    tokens = [
        OCRToken(text="TAX INVOICE # INV-900", xmin=50, ymin=50, xmax=300, ymax=80, confidence=0.99),
        OCRToken(text="ITEM DESCRIPTION", xmin=50, ymin=120, xmax=200, ymax=140, confidence=0.95),
        OCRToken(text="QTY", xmin=260, ymin=120, xmax=290, ymax=140, confidence=0.95),
        OCRToken(text="RATE", xmin=350, ymin=120, xmax=400, ymax=140, confidence=0.95),
        OCRToken(text="TOTAL", xmin=470, ymin=120, xmax=530, ymax=140, confidence=0.95),
        # Line with SKU code 90021, Quantity 5, Unit Rate 20.00, Line Total 100.00
        OCRToken(text="Cable SKU-90021", xmin=50, ymin=170, xmax=190, ymax=190, confidence=0.95),
        OCRToken(text="5", xmin=265, ymin=170, xmax=285, ymax=190, confidence=0.95),
        OCRToken(text="20.00", xmin=355, ymin=170, xmax=395, ymax=190, confidence=0.95),
        OCRToken(text="100.00", xmin=475, ymin=170, xmax=525, ymax=190, confidence=0.95),
        # Totals
        OCRToken(text="TOTAL AMOUNT", xmin=300, ymin=250, xmax=420, ymax=270, confidence=0.95),
        OCRToken(text="100.00", xmin=475, ymin=250, xmax=525, ymax=270, confidence=0.95),
    ]

    parser = SpatialInvoiceParser(tokens)
    items = parser.extract_line_items()
    assert len(items) == 1
    assert items[0].quantity == 5.0
    assert items[0].unit_price == 20.00
    assert items[0].line_total == 100.00
    assert items[0].is_math_sound is True
    assert "Cable" in items[0].description

def test_excel_export_with_none_invoice_number(tmp_path):
    from invoice_audit_engine.schema import InvoiceExtractionResult, InvoiceHeader, LineItem, InvoiceTotals, AuditValidationFlags
    from invoice_audit_engine.excel_exporter import export_audit_to_excel

    # Document where invoice_number is None and notes has None entries
    res_none = InvoiceExtractionResult(
        file_path="C:/invoices/unnamed_receipt.jpg",
        file_hash_sha256="fakehash_none_inv_123",
        document_type="INVOICE",
        is_invoice_archetype=True,
        quality_passed=True,
        quality_metrics={"passed": True, "issues": []},
        header=InvoiceHeader(
            vendor_name="Corner Store",
            invoice_number=None, # None invoice number
            currency="USD"
        ),
        line_items=[
            LineItem(
                description="General Supplies",
                quantity=2.0,
                unit_price=25.0,
                line_total=50.0,
                calculated_total=50.0,
                line_variance=0.0,
                is_math_sound=True
            )
        ],
        totals=InvoiceTotals(
            subtotal=50.0,
            total_amount=50.0
        ),
        audit_flags=AuditValidationFlags(
            is_mathematically_sound=True,
            notes=[]
        ),
        raw_token_count=12,
        average_ocr_confidence=0.92
    )

    out_file = str(tmp_path / "test_none_inv.xlsx")
    saved = export_audit_to_excel([res_none], out_file)
    assert os.path.exists(saved)
    assert os.path.getsize(saved) > 5000

def test_null_fuzzing_matrix(tmp_path):
    from invoice_audit_engine.schema import InvoiceExtractionResult, InvoiceHeader, LineItem, InvoiceTotals, AuditValidationFlags
    from invoice_audit_engine.excel_exporter import export_audit_to_excel

    fuzz_docs = [
        # 1. Total Null State (literally every optional field is None)
        InvoiceExtractionResult(
            file_path="C:/fuzz/total_null.png",
            file_hash_sha256="fuzz_hash_001",
            document_type="INVOICE",
            is_invoice_archetype=True,
            exclusion_reason=None,
            quality_passed=True,
            quality_metrics={"issues": []},
            header=InvoiceHeader(
                vendor_name=None, invoice_number=None, invoice_date=None, due_date=None,
                tax_id=None, po_number=None, document_type="INVOICE", currency=None
            ),
            line_items=[
                LineItem(description="Item 1", quantity=None, unit_price=None, line_total=None,
                         calculated_total=None, line_variance=None, is_math_sound=True, raw_text=None),
                LineItem(description="Item 2 With Symbols !@#$%^&*()_+", quantity=0.0, unit_price=-10.5, line_total=0.0,
                         calculated_total=0.0, line_variance=0.0, is_math_sound=False, raw_text="raw")
            ],
            totals=InvoiceTotals(subtotal=None, discount=None, tax_amount=None, tax_rate_pct=None, total_amount=None, currency=None),
            audit_flags=AuditValidationFlags(
                is_mathematically_sound=False, math_mismatch_totals=True, totals_discrepancy_amount=15.0,
                math_mismatch_line_items=True, lines_discrepancy_amount=25.0, line_item_calc_mismatches_count=1,
                missing_tax_id_on_taxed_invoice=True, low_ocr_confidence=True, notes=[None, "", "Discrepancy detected"]
            ),
            raw_token_count=0,
            average_ocr_confidence=0.0
        ),
        # 2. Rejected document with empty / None issues
        InvoiceExtractionResult(
            file_path="C:/fuzz/rejected_null.jpg",
            file_hash_sha256="fuzz_hash_002",
            document_type="NON_INVOICE_DOCUMENT",
            is_invoice_archetype=False,
            exclusion_reason="Unknown document",
            quality_passed=False,
            quality_metrics={"passed": False, "issues": [None, "Camera blur"]},
            header=None,
            line_items=[],
            totals=None,
            audit_flags=AuditValidationFlags(
                is_mathematically_sound=False,
                notes=[None, "Blur failure"]
            ),
            raw_token_count=0,
            average_ocr_confidence=0.0
        ),
        # 3. Unicode and extreme values
        InvoiceExtractionResult(
            file_path="C:/fuzz/unicode_extreme.tiff",
            file_hash_sha256="fuzz_hash_003",
            document_type="INVOICE",
            is_invoice_archetype=True,
            quality_passed=True,
            quality_metrics={"issues": []},
            header=InvoiceHeader(
                vendor_name="مؤسسة التقنية العالمية 🌟",
                invoice_number="INV/2026/000999/XYZ",
                currency="AED"
            ),
            line_items=[
                LineItem(description="معدات شبكات 🌟", quantity=1000000.0, unit_price=0.001, line_total=1000.0,
                         calculated_total=1000.0, line_variance=0.0, is_math_sound=True)
            ],
            totals=InvoiceTotals(subtotal=1000.0, total_amount=1000.0),
            audit_flags=AuditValidationFlags(is_mathematically_sound=True, notes=[]),
            raw_token_count=50,
            average_ocr_confidence=0.99
        )
    ]

    out_file = str(tmp_path / "fuzz_report.xlsx")
    saved = export_audit_to_excel(fuzz_docs, out_file)
    assert os.path.exists(saved)
    assert os.path.getsize(saved) > 5000

def test_heterogeneous_batch_population(tmp_path):
    from invoice_audit_engine.schema import InvoiceExtractionResult, InvoiceHeader, LineItem, InvoiceTotals, AuditValidationFlags
    from invoice_audit_engine.checkpoint_db import CheckpointLedger
    from invoice_audit_engine.excel_exporter import export_audit_to_excel

    db_path = str(tmp_path / "hetero_ledger.db")
    ledger = CheckpointLedger(db_path)

    docs = [
        # 1. Clean Tax Invoice
        InvoiceExtractionResult(
            file_path="C:/pop/inv_01.pdf", file_hash_sha256="h1", document_type="INVOICE", is_invoice_archetype=True,
            quality_passed=True, quality_metrics={"issues": []},
            header=InvoiceHeader(vendor_name="Acme Supplies", invoice_number="INV-100", currency="USD"),
            line_items=[LineItem(description="Widget A", quantity=2.0, unit_price=50.0, line_total=100.0, calculated_total=100.0, is_math_sound=True)],
            totals=InvoiceTotals(subtotal=100.0, total_amount=100.0),
            audit_flags=AuditValidationFlags(is_mathematically_sound=True, notes=[]),
            raw_token_count=20, average_ocr_confidence=0.95
        ),
        # 2. Excluded Statement of Account
        InvoiceExtractionResult(
            file_path="C:/pop/soa_01.pdf", file_hash_sha256="h2", document_type="STATEMENT_OF_ACCOUNT", is_invoice_archetype=False,
            quality_passed=True, quality_metrics={"issues": []},
            header=InvoiceHeader(vendor_name="Acme Supplies", invoice_number="SOA-888", currency="USD"),
            line_items=[], totals=InvoiceTotals(subtotal=None, total_amount=None),
            audit_flags=AuditValidationFlags(is_mathematically_sound=True, notes=["EXCLUDED_FROM_AUDIT"]),
            raw_token_count=35, average_ocr_confidence=0.90
        ),
        # 3. Cash Slip with None invoice number
        InvoiceExtractionResult(
            file_path="C:/pop/cash_slip.jpg", file_hash_sha256="h3", document_type="INVOICE", is_invoice_archetype=True,
            quality_passed=True, quality_metrics={"issues": []},
            header=InvoiceHeader(vendor_name="Local Hardware", invoice_number=None, currency="USD"),
            line_items=[LineItem(description="Widget A", quantity=5.0, unit_price=60.0, line_total=300.0, calculated_total=300.0, is_math_sound=True)],
            totals=InvoiceTotals(subtotal=300.0, total_amount=300.0),
            audit_flags=AuditValidationFlags(is_mathematically_sound=True, notes=[]),
            raw_token_count=15, average_ocr_confidence=0.88
        ),
        # 4. Rejected Blurry Scan
        InvoiceExtractionResult(
            file_path="C:/pop/blurry.jpg", file_hash_sha256="h4", document_type="INVOICE", is_invoice_archetype=True,
            quality_passed=False, quality_metrics={"passed": False, "issues": ["Blur"]},
            header=None, line_items=[], totals=None,
            audit_flags=AuditValidationFlags(is_mathematically_sound=False, notes=["REJECTED_AT_QUALITY_GATE"]),
            raw_token_count=0, average_ocr_confidence=0.0
        ),
        # 5. Math Mismatch Invoice
        InvoiceExtractionResult(
            file_path="C:/pop/inv_bad_math.png", file_hash_sha256="h5", document_type="INVOICE", is_invoice_archetype=True,
            quality_passed=True, quality_metrics={"issues": []},
            header=InvoiceHeader(vendor_name="Faulty Corp", invoice_number="BAD-01", currency="USD"),
            line_items=[LineItem(description="Service Fee", quantity=1.0, unit_price=100.0, line_total=100.0, calculated_total=100.0, is_math_sound=True)],
            totals=InvoiceTotals(subtotal=100.0, tax_amount=0.0, total_amount=150.0),
            audit_flags=AuditValidationFlags(is_mathematically_sound=False, math_mismatch_totals=True, totals_discrepancy_amount=50.0, notes=["Totals mismatch"]),
            raw_token_count=25, average_ocr_confidence=0.91
        )
    ]

    for d in docs:
        ledger.save_result(d)

    loaded = ledger.load_all_results()
    assert len(loaded) == 5

    out_file = str(tmp_path / "hetero_master_report.xlsx")
    saved = export_audit_to_excel(loaded, out_file)
    assert os.path.exists(saved)

    import openpyxl
    wb = openpyxl.load_workbook(saved)
    ws_kpi = wb["Executive_Summary"]
    kpis = {ws_kpi.cell(row=r, column=1).value: ws_kpi.cell(row=r, column=2).value for r in range(5, 15)}
    assert kpis["Total Documents Examined"] == 5
    assert kpis["Clean Invoices Audited"] == 3
    assert kpis["Non-Invoice Documents Filtered & Excluded"] == 1
    assert kpis["Documents Rejected at Ingestion Gate"] == 1
    assert kpis["Mathematical Reconciliation Failures (Totals)"] == 1
    assert round(kpis["Total Billed Spend Audited"], 2) == 550.00

    # Tab 4: Verify Price Variance Analytics detected variance on 'Widget A' (50.00 vs 60.00)
    ws_price = wb["Price_Variance_Analytics"]
    assert ws_price.cell(row=2, column=1).value == "Widget A"
    assert ws_price.cell(row=2, column=3).value == 50.0
    assert ws_price.cell(row=2, column=4).value == 60.0
    assert ws_price.cell(row=2, column=9).value == "ALERT: PRICE VARIANCE"
    assert "INV-100" in str(ws_price.cell(row=2, column=10).value)


